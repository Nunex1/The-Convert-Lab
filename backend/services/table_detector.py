import re
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from backend.services.errors import ConversionError


def typed_value(raw):
    if raw is None:
        return None, None
    value = re.sub(r'\s+', ' ', str(raw)).strip()
    if not value:
        return None, None
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(value, fmt), 'dd/mm/yyyy'
        except ValueError:
            pass
    # Preserve identifiers and leading zeros. A lone 1.234 is ambiguous; keep as text.
    numeric = value.replace('R$', '').replace('$', '').replace('€', '').replace(' ', '')
    percent = numeric.endswith('%')
    numeric = numeric.rstrip('%')
    negative = numeric.startswith('(') and numeric.endswith(')')
    numeric = numeric.strip('()')
    if re.match(r'^[-+]?0\d', numeric) or re.fullmatch(r'[-+]?\d{1,3}\.\d{3}', numeric):
        return value, None
    if re.fullmatch(r'[-+]?\d{1,3}(\.\d{3})+,\d+|[-+]?\d+,\d+', numeric):
        numeric = numeric.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'[-+]?\d{1,3}(,\d{3})+(\.\d+)?', numeric):
        numeric = numeric.replace(',', '')
    if re.fullmatch(r'[-+]?\d+(\.\d+)?', numeric) and len(numeric.replace('.', '').lstrip('+-')) <= 15:
        number = float(numeric) if '.' in numeric or percent else int(numeric)
        number = -abs(number) if negative else number
        if percent:
            return number / 100, '0.00%'
        if 'R$' in value:
            return number, '"R$" #,##0.00;[Red]-"R$" #,##0.00'
        return number, '#,##0.00' if isinstance(number, float) else '#,##0'
    return value, None


def write_tables(tables: list[tuple[str, list]], target: Path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    total_cells = 0
    for name, rows in tables:
        rows = [row for row in rows if any(str(c or '').strip() for c in row)]
        if not rows:
            continue
        width = max(map(len, rows))
        keep = [i for i in range(width) if any(i < len(r) and str(r[i] or '').strip() for r in rows)]
        total_cells += len(rows) * len(keep)
        if total_cells > 500_000 or len(keep) > 16384 or len(rows) > 1_048_576:
            raise ConversionError('As tabelas excedem o limite de processamento. Divida o documento.')
        title = re.sub(r'[\\/*?:\[\]]', '', name).strip(" '")[:31] or 'Tabela'
        sheet = workbook.create_sheet(title)
        for rindex, row in enumerate(rows, 1):
            for cindex, original in enumerate(keep, 1):
                value, fmt = typed_value(row[original] if original < len(row) else None)
                cell = sheet.cell(rindex, cindex, value)
                if isinstance(value, str):
                    cell.data_type = 's'  # Never turn untrusted document text into spreadsheet formulas.
                if fmt:
                    cell.number_format = fmt
                if rindex == 1:
                    cell.font = Font(color='FFFFFF', bold=True)
                    cell.fill = PatternFill('solid', fgColor='263647')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(14, max(len(str(c.value or '')) for c in column) + 3))
    if not workbook.sheetnames:
        raise ConversionError('Não foi possível identificar tabelas neste documento.')
    workbook.save(target)
