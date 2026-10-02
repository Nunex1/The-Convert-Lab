from pathlib import Path
from backend.services.errors import ConversionError


def validate_result(path: Path, extension: str) -> None:
    if not path.exists() or path.stat().st_size == 0:
        raise ConversionError('A conversão não produziu conteúdo utilizável.')
    if path.stat().st_size > 150 * 1024 * 1024:
        raise ConversionError('O resultado excede o limite de 150 MB. Divida o documento e tente novamente.')
    if extension == 'md':
        if not path.read_text(encoding='utf-8').strip():
            raise ConversionError('Não foi possível extrair texto deste documento.')
    elif extension == 'docx':
        from docx import Document
        Document(path)
    elif extension == 'xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True)
        try:
            if not workbook.sheetnames or not any(s.max_row >= 1 for s in workbook):
                raise ConversionError('Não foram encontradas tabelas utilizáveis.')
        finally:
            workbook.close()
    elif extension == 'pptx':
        from pptx import Presentation
        if not Presentation(path).slides:
            raise ConversionError('Não foi possível gerar slides deste documento.')
