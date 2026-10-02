from io import BytesIO
from pathlib import Path
import pymupdf
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from backend.services.layout_detector import editable_page


def convert(source: Path, target: Path) -> list[str]:
    presentation = Presentation()
    fallback_pages = 0
    with pymupdf.open(source) as document:
        first = document[0].rect
        scale = min(1, 1440 / max(first.width, first.height))
        width, height = max(72, first.width * scale), max(72, first.height * scale)
        presentation.slide_width, presentation.slide_height = Pt(width), Pt(height)
        for page in document:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            blocks = page.get_text('dict')['blocks']
            factor = min(width / page.rect.width, height / page.rect.height)
            offset_x = (width - page.rect.width * factor) / 2
            offset_y = (height - page.rect.height * factor) / 2
            if editable_page(page, blocks):
                for block in blocks:
                    if block['type'] != 0:
                        continue
                    for line in block['lines']:
                        for span in line['spans']:
                            x0, y0, x1, y1 = span['bbox']
                            shape = slide.shapes.add_textbox(Pt(offset_x + x0 * factor), Pt(offset_y + y0 * factor), Pt(max(2, (x1 - x0 + 6) * factor)), Pt(max(2, (y1 - y0 + 4) * factor)))
                            frame = shape.text_frame
                            frame.margin_left = frame.margin_top = frame.margin_right = frame.margin_bottom = 0
                            frame.word_wrap = False
                            run = frame.paragraphs[0].add_run()
                            run.text = span['text']
                            run.font.name = 'Arial'
                            run.font.size = Pt(max(1, span['size'] * factor))
                            run.font.bold = bool(span['flags'] & 16)
                            run.font.italic = bool(span['flags'] & 2)
                            color = span['color']
                            run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
            else:
                fallback_pages += 1
                zoom = min(2, 2400 / max(page.rect.width, page.rect.height))
                image = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).tobytes('png')
                slide.shapes.add_picture(BytesIO(image), Pt(offset_x), Pt(offset_y), width=Pt(page.rect.width * factor), height=Pt(page.rect.height * factor))
    presentation.save(target)
    return [f'{fallback_pages} página(s) inserida(s) como imagem para preservar o layout. Essas páginas não têm objetos editáveis.'] if fallback_pages else []
