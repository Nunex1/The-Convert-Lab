from io import BytesIO
from pathlib import Path
import pymupdf
from docx import Document
from docx.shared import Inches
from pdf2docx import Converter


def fallback(document, target):
    output = Document()
    for index, page in enumerate(document):
        if index:
            output.add_page_break()
        blocks = page.get_text('dict', sort=True)['blocks']
        for block in blocks:
            if block['type'] == 0:
                for line in block['lines']:
                    paragraph = output.add_paragraph()
                    for span in line['spans']:
                        run = paragraph.add_run(span['text'])
                        run.bold, run.italic = bool(span['flags'] & 16), bool(span['flags'] & 2)
        # A page image retains diagrams and scanned content when reconstruction fails.
        if page.get_images() or page.get_drawings() or not page.get_text().strip():
            zoom = min(1.5, 2000 / max(page.rect.width, page.rect.height))
            image = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).tobytes('png')
            output.add_picture(BytesIO(image), width=Inches(6))
    output.save(target)


def convert(source: Path, target: Path) -> list[str]:
    with pymupdf.open(source) as document:
        scanned = any(len(p.get_text().strip()) < 12 and p.get_images() for p in document)
        if not scanned:
            converter = None
            try:
                converter = Converter(str(source))
                converter.convert(str(target), multi_processing=False)
                result = Document(target)
                expected = sum(len(p.get_text().split()) for p in document)
                actual = len(' '.join(p.text for p in result.paragraphs).split()) + sum(len(c.text.split()) for t in result.tables for row in t.rows for c in row.cells)
                if expected and actual < expected * .65:
                    raise ValueError('Incomplete text reconstruction')
                return ['Layouts complexos podem apresentar diferenças de espaçamento e tipografia.']
            except Exception:
                pass
            finally:
                if converter:
                    converter.close()
        fallback(document, target)
    return ['Foi utilizado um layout simplificado: texto disponível editável e imagens das páginas para preservar conteúdo visual. Páginas digitalizadas permanecem como imagens; OCR não está disponível.']
