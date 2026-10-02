import os
from pathlib import Path
import tempfile
os.environ['TEMP_DIR'] = tempfile.mkdtemp(prefix='convertlab-test-')

import pytest
import pymupdf
from fastapi.testclient import TestClient
from backend.main import app


@pytest.fixture
def client():
    with TestClient(app) as client:
        yield client


@pytest.fixture
def pdf(tmp_path):
    path = tmp_path / 'Relatorio.pdf'
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 65), 'Relatorio financeiro', fontsize=22)
    page.insert_text((50, 100), 'Dados de exemplo para validar a conversao real.', fontsize=11)
    for y in (140, 170, 200, 230):
        page.draw_line((50, y), (500, y))
    for x in (50, 275, 500):
        page.draw_line((x, 140), (x, 230))
    for y, a, b in [(160, 'Item', 'Valor'), (190, 'Receita', '1250,50'), (220, 'Custo', '500,00')]:
        page.insert_text((60, y), a, fontsize=11)
        page.insert_text((285, y), b, fontsize=11)
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def docx(tmp_path):
    from docx import Document
    path = tmp_path / 'Documento.docx'
    doc = Document()
    doc.add_heading('Planejamento', 0)
    doc.add_heading('Objetivos', 1)
    doc.add_paragraph('Conteudo editavel e verificavel.').add_run(' Destaque').bold = True
    doc.add_paragraph('Primeiro item', style='List Bullet')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Projeto'
    table.cell(0, 1).text = 'Status'
    table.cell(1, 0).text = 'ConvertLab'
    table.cell(1, 1).text = 'Ativo'
    doc.save(path)
    return path


@pytest.fixture
def mht(tmp_path):
    from email.message import EmailMessage
    message = EmailMessage()
    message.set_type('multipart/related')
    html = EmailMessage()
    html.set_content('<html><table><caption>Financeiro</caption><tr><th>Item</th><th>Valor</th><th>Data</th></tr><tr><td>Receita</td><td>R$ 1.250,50</td><td>01/10/2026</td></tr><tr><td>=HYPERLINK("https://invalid")</td><td>25%</td><td>02/10/2026</td></tr></table></html>', subtype='html', charset='utf-8')
    message.attach(html)
    path = tmp_path / 'Tabelas.mht'
    path.write_bytes(message.as_bytes())
    return path
