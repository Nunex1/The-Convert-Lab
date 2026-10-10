from io import BytesIO
import json
from urllib.parse import unquote
import pymupdf
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from backend.tests.test_initial_flows import convert


def test_pdf_word(client, pdf):
    response = convert(client, pdf, 'docx', 'application/pdf')
    assert response.status_code == 200, response.text
    document = Document(BytesIO(response.content))
    text = '\n'.join(p.text for p in document.paragraphs)
    text += '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert 'Relatorio financeiro' in text
    assert '1250' in text


def test_pdf_excel(client, pdf):
    response = convert(client, pdf, 'xlsx', 'application/pdf')
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet['A2'].value == 'Receita'
    assert sheet['B2'].value == 1250.5


def test_pdf_pptx_visual_fallback(client, pdf):
    response = convert(client, pdf, 'pptx', 'application/pdf')
    assert response.status_code == 200, response.text
    presentation = Presentation(BytesIO(response.content))
    assert len(presentation.slides) == 1
    assert any(s.shape_type == 13 for s in presentation.slides[0].shapes)
    assert 'imagem' in json.loads(unquote(response.headers['X-Conversion-Warnings']))[0]


def test_pdf_pptx_editable(client, tmp_path):
    path = tmp_path / 'text.pdf'
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((70, 90), 'Texto editavel', fontsize=20)
        document.save(path)
    response = convert(client, path, 'pptx', 'application/pdf')
    assert response.status_code == 200, response.text
    presentation = Presentation(BytesIO(response.content))
    assert 'Texto editavel' in ' '.join(s.text for s in presentation.slides[0].shapes if s.has_text_frame)


def test_mhtml(client, mht):
    path = mht.with_suffix('.mhtml')
    path.write_bytes(mht.read_bytes())
    response = convert(client, path, 'xlsx', 'application/octet-stream')
    assert response.status_code == 200
    assert load_workbook(BytesIO(response.content)).active['B2'].value == 1250.5


def test_pptx_mixed_pages_preserve_content_and_proportions(client, tmp_path):
    path = tmp_path / 'mixed.pdf'
    with pymupdf.open() as document:
        page = document.new_page(width=600, height=400)
        page.insert_text((40, 60), 'Editable first page', fontsize=20)
        page = document.new_page(width=400, height=800)
        page.insert_text((40, 60), 'Vector layout')
        page.draw_rect(pymupdf.Rect(20, 80, 350, 750), color=(1, 0, 0))
        page = document.new_page(width=800, height=400)
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 120, 80), False)
        pixmap.clear_with(180)
        page.insert_image(pymupdf.Rect(0, 0, 800, 400), stream=pixmap.tobytes('png'))
        page = document.new_page(width=600, height=400)
        page.insert_text((40, 60), 'Rotated page')
        page.set_rotation(90)
        document.save(path)
    response = convert(client, path, 'pptx', 'application/pdf')
    assert response.status_code == 200, response.text
    presentation = Presentation(BytesIO(response.content))
    assert len(presentation.slides) == 4
    assert 'Editable first page' in ' '.join(s.text for s in presentation.slides[0].shapes if s.has_text_frame)
    for slide, ratio in zip(list(presentation.slides)[1:], [.5, 2, 400 / 600]):
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        assert len(pictures) == 1
        picture = pictures[0]
        assert abs(picture.width / picture.height - ratio) < .001
        assert picture.left >= 0 and picture.top >= 0
        assert picture.left + picture.width <= presentation.slide_width + 1
        assert picture.top + picture.height <= presentation.slide_height + 1
        assert picture.crop_top == picture.crop_bottom == picture.crop_left == picture.crop_right == 0
        with pymupdf.open(stream=picture.image.blob, filetype='png') as image:
            assert len(image) == 1
    warnings = json.loads(unquote(response.headers['X-Conversion-Warnings']))
    assert '3 página(s)' in warnings[0]


def test_scanned_pdf_is_honest_about_ocr(client, tmp_path):
    path = tmp_path / 'scan.pdf'
    with pymupdf.open() as document:
        page = document.new_page()
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 100, 100), False)
        pixmap.clear_with(128)
        page.insert_image(page.rect, stream=pixmap.tobytes('png'))
        document.save(path)
    for output in ('md', 'xlsx'):
        response = convert(client, path, output, 'application/pdf')
        assert response.status_code == 422 and 'OCR' in response.json()['detail']
    response = convert(client, path, 'docx', 'application/pdf')
    assert response.status_code == 200
    assert len(Document(BytesIO(response.content)).inline_shapes) == 1
