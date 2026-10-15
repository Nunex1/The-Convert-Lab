import csv
from datetime import datetime
import io
import json
from pathlib import Path
from urllib.parse import unquote
import zipfile

import pytest
from openpyxl import Workbook

from backend.config import TEMP_ROOT
from backend.converters.xlsx.xlsx_to_csv import convert, normalize_options
from backend.converters.xlsx.xlsx_to_xlsb import licensed_engine
from backend.services.errors import ConversionError
from backend.services.file_validator import validate_file
from backend.services.result_validator import validate_result
from backend.services.capabilities import available_formats


@pytest.fixture
def xlsx(tmp_path):
    path = tmp_path / 'Relatório.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'São Paulo'
    ws.append(['ação; vírgula, "aspas"\nlinha', 12.25, True, None, '00123', datetime(2026, 10, 14, 12, 30), '=1+2', '@soma', -42])
    ws['J1'] = 123
    ws['J1'].number_format = '00000'
    wb.create_sheet('Outra').append(['outra aba', 99])
    wb.save(path)
    return path


@pytest.mark.parametrize('delimiter,bom', [(';', False), (',', True)])
def test_real_csv_http(client, xlsx, delimiter, bom):
    response = client.post('/api/convert', files={'file': (xlsx.name, xlsx.read_bytes(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
                           data={'output_format': 'csv', 'csv_options': json.dumps({'delimiter': delimiter, 'bom': bom})})
    assert response.status_code == 200, response.text
    assert response.content.startswith(b'\xef\xbb\xbf') == bom
    rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'), newline=''), delimiter=delimiter))
    assert rows == [['ação; vírgula, "aspas"\nlinha', '12.25', 'TRUE', '', '00123', '2026-10-14T12:30:00', '', "'@soma", '-42', '00123']]
    warnings = json.loads(unquote(response.headers['x-conversion-warnings']))
    assert 'São Paulo' in warnings[0] and '1 outra' in warnings[0]
    assert any('sem cache' in w for w in warnings)
    assert any('apóstrofo' in w for w in warnings)
    assert 'Relatório.csv' in unquote(response.headers['content-disposition'])
    assert response.headers['content-type'] == 'text/csv; charset=utf-8'
    assert 'input_validation;dur=' in response.headers['server-timing']
    assert not list(TEMP_ROOT.glob('job-*'))


def test_sheet_and_exact_text(xlsx, tmp_path):
    target = tmp_path / 'out.csv'
    warnings = convert(xlsx, target, {'sheet': 'Outra'})
    assert target.read_text().strip() == 'outra aba;99'
    assert 'Outra' in warnings[0]
    warnings = convert(xlsx, target, {'protect_formulas': False})
    assert ';@soma;' in target.read_text()
    assert any('escolha explícita' in w for w in warnings)
    with pytest.raises(ConversionError, match='não existe'):
        convert(xlsx, target, {'sheet': 'missing'})


@pytest.mark.parametrize('options', [[], 'bad', {'delimiter': '|'}, {'bom': 'yes'}, {'protect_formulas': 1}, {'sheet': 'x' * 32}, {'unknown': True}])
def test_options_rejected(options):
    with pytest.raises(ConversionError):
        normalize_options(options)


@pytest.mark.parametrize('payload', ['{', '[]', '{"delimiter":"|"}', '{"bom":"false"}'])
def test_invalid_options_http(client, xlsx, payload):
    response = client.post('/api/convert', files={'file': (xlsx.name, xlsx.read_bytes())}, data={'output_format': 'csv', 'csv_options': payload})
    assert response.status_code == 400
    assert not list(TEMP_ROOT.glob('job-*'))


def rewrite(path, name, contents):
    with zipfile.ZipFile(path) as archive:
        parts = {i.filename: archive.read(i) for i in archive.infolist()}
    parts[name] = contents
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for key, value in parts.items():
            archive.writestr(key, value)


@pytest.mark.parametrize('name,contents', [
    ('../bad.xml', '<a/>'), ('xl/vbaProject.bin', 'macro'),
    ('xl/externalLinks/link.xml', '<a/>'), ('xl/connections.xml', '<a/>'),
    ('xl/evil.xml', '<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>'),
    ('xl/_rels/evil.rels', '<Relationships><Relationship TargetMode="External" Target="https://invalid"/></Relationships>'),
    ('xl/evil.xml', '<root><sheetProtection/></root>'),
    ('xl/huge.bin', b'0' * 2_000_000),
], ids=['path', 'macro', 'external-part', 'connection', 'entity', 'external-relation', 'protected', 'zip-ratio'])
def test_unsafe_xlsx(xlsx, name, contents):
    rewrite(xlsx, name, contents)
    with pytest.raises(ConversionError):
        validate_file(xlsx, 'xlsx', 'application/zip')


def test_corrupt_and_renamed_xlsx(tmp_path):
    source = tmp_path / 'bad.xlsx'
    source.write_bytes(b'not a zip')
    with pytest.raises(ConversionError):
        validate_file(source, 'xlsx', None)


@pytest.mark.parametrize('formula', ['=WEBSERVICE("https://example.invalid")', '=HYPERLINK("https://example.invalid")', '=RTD("server",,"topic")', "='[external.xlsx]Sheet1'!A1", "=cmd|' /C test'!A0"])
def test_external_formula_rejected(xlsx, formula):
    from openpyxl import load_workbook
    workbook = load_workbook(xlsx)
    workbook.active['A1'] = formula
    workbook.save(xlsx)
    with pytest.raises(ConversionError, match='fórmulas'):
        validate_file(xlsx, 'xlsx', None)


def test_formula_literal_is_not_external(xlsx):
    from openpyxl import load_workbook
    workbook = load_workbook(xlsx)
    workbook.active['A1'] = '=CONCAT("HYPERLINK(","texto")'
    workbook.save(xlsx)
    validate_file(xlsx, 'xlsx', None)


def test_out_of_bounds_coordinates(xlsx):
    with zipfile.ZipFile(xlsx) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').replace(b'r="A1"', b'r="ZZZ9999999"')
    rewrite(xlsx, 'xl/worksheets/sheet1.xml', xml)
    with pytest.raises(ConversionError, match='coordenadas'):
        validate_file(xlsx, 'xlsx', None)


@pytest.mark.parametrize('bom', [True, False])
def test_empty_sheet(tmp_path, bom):
    source, target = tmp_path / 'empty.xlsx', tmp_path / 'empty.csv'
    Workbook().save(source)
    convert(source, target, {'bom': bom})
    validate_result(target, 'csv')
    assert target.read_bytes() == (b'\xef\xbb\xbf' if bom else b'') + b'\r\n'


@pytest.mark.parametrize('rows', [100, 10000, 100000])
def test_large_csv_values_and_row_count(tmp_path, rows):
    source, target = tmp_path / 'large.xlsx', tmp_path / 'large.csv'
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Dados')
    for index in range(rows):
        sheet.append([index, f'ação {index}; "teste"\nlinha', '00123', True, 12.25])
    workbook.save(source)
    validate_file(source, 'xlsx', None)
    convert(source, target)
    validate_result(target, 'csv')
    with target.open(encoding='utf-8', newline='') as stream:
        count = 0
        for count, record in enumerate(csv.reader(stream, delimiter=';'), 1):
            index = count - 1
            assert record == [str(index), f'ação {index}; "teste"\nlinha', '00123', 'TRUE', '12.25']
        assert count == rows


def test_duplicate_zip_entries_rejected(xlsx):
    with zipfile.ZipFile(xlsx, 'a') as archive:
        with pytest.warns(UserWarning):
            archive.writestr('xl/workbook.xml', '<a/>')
    with pytest.raises(ConversionError, match='duplicadas'):
        validate_file(xlsx, 'xlsx', None)


def test_renaming_xlsx_is_not_xlsb(xlsx):
    pytest.importorskip('pyxlsb')
    from backend.converters.xlsx.xlsx_to_xlsb import validate_xlsb
    with pytest.raises(ConversionError, match='autêntica'):
        validate_xlsb(xlsx)


def test_licensed_xlsb_roundtrip(xlsx, tmp_path):
    import os
    if not os.getenv('ASPOSE_CELLS_LICENSE_PATH') or os.getenv('ENABLE_XLSB') != '1':
        pytest.skip('Requires an explicitly enabled, valid commercial license')
    from backend.converters.xlsx.xlsx_to_xlsb import convert as to_xlsb
    from pyxlsb import open_workbook
    target = tmp_path / 'real.xlsb'
    to_xlsb(xlsx, target)
    validate_result(target, 'xlsb')
    with open_workbook(str(target)) as result:
        assert result.sheets == ['São Paulo', 'Outra']
        with result.get_sheet('Outra') as sheet:
            assert [c.v for c in next(sheet.rows())] == ['outra aba', 99]


def test_missing_formula_cache_and_stale_dimension(xlsx, tmp_path):
    with zipfile.ZipFile(xlsx) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').replace(b'<f>1+2</f><v></v>', b'<f>1+2</f><v>3</v>').replace(b'ref="A1:J1"', b'ref="A1:A1"')
    rewrite(xlsx, 'xl/worksheets/sheet1.xml', xml)
    target = tmp_path / 'out.csv'
    convert(xlsx, target)
    with target.open(newline='', encoding='utf-8') as stream:
        assert next(csv.reader(stream, delimiter=';'))[6] == '3'


def test_xlsb_disabled_without_license(monkeypatch, client, xlsx):
    monkeypatch.delenv('ENABLE_XLSB', raising=False)
    available_formats.cache_clear()
    assert client.get('/api/formats').json()['xlsx'] == ['csv']
    with pytest.raises(ConversionError):
        licensed_engine()
    response = client.post('/api/convert', files={'file': (xlsx.name, xlsx.read_bytes())}, data={'output_format': 'xlsb'})
    assert response.status_code == 400


def test_invalid_license_fails_closed(monkeypatch):
    monkeypatch.setenv('ENABLE_XLSB', '1')
    monkeypatch.setenv('ASPOSE_CELLS_LICENSE_PATH', '/nonexistent/license.lic')
    available_formats.cache_clear()
    try:
        assert 'xlsb' not in available_formats()['xlsx']
        with pytest.raises(ConversionError):
            licensed_engine()
    finally:
        available_formats.cache_clear()


def test_csv_validator_rejects_corruption(tmp_path):
    path = tmp_path / 'bad.csv'
    path.write_text('"unclosed', encoding='utf-8')
    with pytest.raises(csv.Error):
        validate_result(path, 'csv')


@pytest.mark.parametrize('text', ['=cmd()', '+cmd()', '-cmd()', '@cmd()', '\tcmd()', '\rstart', '\nstart', '  =cmd()'])
def test_formula_like_text_protected(tmp_path, text):
    source, target = tmp_path / 'source.xlsx', tmp_path / 'out.csv'
    wb = Workbook()
    wb.active['A1'] = text
    wb.active['A1'].data_type = 's'
    wb.save(source)
    convert(source, target)
    with target.open(encoding='utf-8', newline='') as stream:
        assert next(csv.reader(stream, delimiter=';')) == ["'" + text]
