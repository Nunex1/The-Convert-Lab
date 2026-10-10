"""Real HTTP conversions with synthetic inputs; run against local or deployed API.

python verification/smoke.py https://the-convert-lab.onrender.com verification/local/production
"""
import argparse
from email.message import EmailMessage
import json
from pathlib import Path
from urllib.parse import unquote

import httpx
import pymupdf
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def samples(root):
    root.mkdir(parents=True, exist_ok=True)
    pdf = root / 'validation.pdf'
    with pymupdf.open() as document:
        page = document.new_page(width=600, height=400)
        page.insert_text((40, 50), 'ConvertLab validation', fontsize=20)
        page = document.new_page(width=400, height=600)
        page.insert_text((40, 50), 'Table validation', fontsize=18)
        for y in (100, 140, 180, 220):
            page.draw_line((40, y), (360, y))
        for x in (40, 200, 360):
            page.draw_line((x, 100), (x, 220))
        for y, left, right in [(125, 'Item', 'Value'), (165, 'Revenue', '1250,50'), (205, 'Cost', '500,00')]:
            page.insert_text((50, y), left)
            page.insert_text((210, y), right)
        document.save(pdf)
    docx = root / 'validation.docx'
    document = Document()
    document.add_heading('ConvertLab validation', 1)
    document.add_paragraph('Editable synthetic content.').add_run(' Emphasis').bold = True
    document.save(docx)
    message = EmailMessage()
    message.set_type('multipart/related')
    html = EmailMessage()
    html.set_content('<table><tr><th>Item</th><th>Value</th></tr><tr><td>Revenue</td><td>1250,50</td></tr></table>', subtype='html')
    message.attach(html)
    for extension in ('mht', 'mhtml'):
        (root / ('validation.' + extension)).write_bytes(message.as_bytes())


def check(path, source, output):
    if output == 'md':
        assert 'ConvertLab validation' in path.read_text(encoding='utf-8')
    elif output == 'docx':
        document = Document(path)
        assert 'ConvertLab validation' in '\n'.join(p.text for p in document.paragraphs)
    elif output == 'xlsx':
        workbook = load_workbook(path, read_only=True)
        try:
            values = [cell.value for sheet in workbook for row in sheet for cell in row]
            assert 'Revenue' in values and 1250.5 in values
        finally:
            workbook.close()
    else:
        presentation = Presentation(path)
        assert len(presentation.slides) == 2
        assert 'ConvertLab validation' in ' '.join(s.text for s in presentation.slides[0].shapes if s.has_text_frame)
        assert any(s.shape_type == 13 for s in presentation.slides[1].shapes)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('url')
    parser.add_argument('directory', type=Path)
    args = parser.parse_args()
    samples(args.directory / 'inputs')
    evidence = {'url': args.url, 'endpoints': {}, 'conversions': []}
    expected_mimes = {
        'md': 'text/markdown', 'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    }
    with httpx.Client(base_url=args.url, timeout=180, headers={'Origin': 'https://theconvertlab.netlify.app'}) as client:
        for endpoint in ('health', 'formats', 'config'):
            response = client.get('/api/' + endpoint)
            response.raise_for_status()
            evidence['endpoints'][endpoint] = response.json()
        for source, outputs in evidence['endpoints']['formats'].items():
            mime = {'pdf': 'application/pdf', 'docx': expected_mimes['docx'], 'mht': 'message/rfc822', 'mhtml': 'application/octet-stream'}[source]
            for output in outputs:
                path = args.directory / 'inputs' / ('validation.' + source)
                response = client.post('/api/convert', files={'file': (path.name, path.read_bytes(), mime)}, data={'output_format': output})
                row = {'input': source, 'output': output, 'http': response.status_code, 'bytes': len(response.content), 'seconds': round(response.elapsed.total_seconds(), 2)}
                if response.status_code == 200:
                    target = args.directory / (source + '-to-' + output + '.' + output)
                    target.write_bytes(response.content)
                    check(target, source, output)
                    assert response.headers['content-type'].startswith(expected_mimes[output])
                    assert 'validation.' + output in unquote(response.headers['content-disposition'])
                    assert response.headers['cache-control'] == 'no-store'
                    assert response.headers['x-content-type-options'] == 'nosniff'
                    assert response.headers['access-control-allow-origin'] == 'https://theconvertlab.netlify.app'
                    row['warnings'] = json.loads(unquote(response.headers['x-conversion-warnings']))
                    row['validated'] = True
                else:
                    row['error'] = response.json()
                    row['validated'] = False
                evidence['conversions'].append(row)
                print(json.dumps(row, ensure_ascii=True), flush=True)
    (args.directory / 'evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if all(row['validated'] for row in evidence['conversions']) else 1


if __name__ == '__main__':
    raise SystemExit(main())
