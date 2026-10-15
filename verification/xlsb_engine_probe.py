"""Evaluation-only compatibility probe. Deletes all watermarked output; never an API conversion."""
import json
from pathlib import Path
import platform
import tempfile
import time
import zipfile

from aspose.cells import Workbook, SaveFormat
from openpyxl import Workbook as SourceWorkbook
from pyxlsb import open_workbook

with tempfile.TemporaryDirectory(prefix='convertlab-evaluation-') as folder:
    source, target = Path(folder) / 'source.xlsx', Path(folder) / 'trial.xlsb'
    wb = SourceWorkbook()
    wb.active.title = 'Dados'
    wb.active.append(['ação', 12.5, True, '00123', '=B1*2'])
    wb.active.merge_cells('A2:B2')
    wb.active['A2'] = 'Merged'
    wb.create_sheet('Outra').append(['segunda'])
    wb.save(source)
    start = time.perf_counter()
    engine = Workbook(str(source))
    licensed = engine.is_licensed
    engine.save(str(target), SaveFormat.XLSB)
    duration = time.perf_counter() - start
    with zipfile.ZipFile(target) as archive:
        binary = 'xl/workbook.bin' in archive.namelist() and 'xl/workbook.xml' not in archive.namelist()
    with open_workbook(str(target)) as result:
        sheets = result.sheets
        with result.get_sheet('Dados') as sheet:
            first = [cell.v for cell in next(sheet.rows())]
    reopened = Workbook(str(target))
    assert binary and first[:4] == ['ação', 12.5, True, '00123']
    assert next(s for s in reopened.worksheets if s.name == 'Dados').cells.get('E1').formula == '=B1*2'
    assert len(sheets) > 2 if not licensed else len(sheets) == 2
    print(json.dumps({'python': platform.python_version(), 'platform': platform.platform(),
                      'licensed': licensed, 'authentic_binary': binary, 'independent_reader': 'pyxlsb 1.0.10',
                      'values_match': True, 'formula_preserved_engine_reopen': True,
                      'source_sheets': 2, 'result_sheets': len(sheets), 'seconds': duration,
                      'bytes': target.stat().st_size,
                      'production_acceptable': licensed, 'output_deleted': True}, indent=2))
