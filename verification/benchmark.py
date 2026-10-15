"""Synthetic, repeatable worker benchmarks; no user documents or production traffic.

Run from repo root: python -m verification.benchmark --output verification/before.json
"""
import argparse
import json
import math
import os
from pathlib import Path
import platform
import shutil
import statistics
import subprocess
import sys
import tempfile
import time
import uuid

import psutil


def measure(command, env, cwd=None):
    started = time.perf_counter()
    peak = 0
    with subprocess.Popen(command, env=env, cwd=cwd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE) as process:
        watched = psutil.Process(process.pid)
        while process.poll() is None:
            if time.perf_counter() - started > 180:
                for child in watched.children(recursive=True):
                    child.kill()
                process.kill()
                raise TimeoutError('Benchmark worker exceeded 180 seconds')
            try:
                peak = max(peak, sum(p.memory_info().rss for p in [watched, *watched.children(recursive=True)]))
            except psutil.NoSuchProcess:
                pass
            time.sleep(.005)
        if process.returncode:
            raise RuntimeError(process.stderr.read().decode(errors='replace'))
    return {'seconds': time.perf_counter() - started, 'peak_rss_mb': peak / 1024**2}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', required=True)
    parser.add_argument('--runs', type=int, default=10)
    parser.add_argument('--xlsx', action='store_true')
    args = parser.parse_args()
    from docx import Document
    import pymupdf
    with tempfile.TemporaryDirectory(prefix='convertlab-bench-') as root:
        root = Path(root)
        env = dict(os.environ, TEMP_DIR=str(root))
        doc = Document()
        for i in range(100):
            doc.add_paragraph(f'Linha {i}: informação sintética para medição.')
        doc.save(root / 'sample.docx')
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((50, 50), 'Benchmark synthetic document')
            pdf.save(root / 'sample.pdf')
        cases = [('docx', 'md', 'sample.docx'), ('pdf', 'pptx', 'sample.pdf')]
        if args.xlsx:
            from openpyxl import Workbook
            for rows in (100, 10000, 100000):
                wb = Workbook(write_only=True)
                ws = wb.create_sheet('Dados')
                for i in range(rows):
                    ws.append([i, f'Texto {i}', 12.34, True, '00123', 'A;B, "C"\nD'])
                name = f'{rows}.xlsx'
                wb.save(root / name)
                cases.append(('xlsx', 'csv', name))
        report = {'python': platform.python_version(), 'platform': platform.platform(),
                  'sampling_ms': 5, 'runs': args.runs, 'cases': {}}
        for label, command in [
            ('python_start', [sys.executable, '-c', 'pass']),
            ('worker_import', [sys.executable, '-c', 'import backend.worker']),
        ]:
            report['cases'][label] = [measure(command, env) for _ in range(args.runs)]
        for source, output, sample in cases:
            observations = []
            for _ in range(args.runs):
                job = root / ('job-' + uuid.uuid4().hex)
                job.mkdir()
                shutil.copyfile(root / sample, job / ('input.' + source))
                (job / 'request.json').write_text(json.dumps({'input': source, 'output': output}))
                observed = measure([sys.executable, '-m', 'backend.worker', str(job)], env)
                observed['output_bytes'] = (job / ('result.' + output)).stat().st_size
                observed['stages'] = json.loads((job / 'status.json').read_text(encoding='utf-8')).get('timings_ms')
                observations.append(observed)
                shutil.rmtree(job)
            report['cases'][sample + '->' + output] = observations
        report['summary'] = {name: {
            'median_seconds': statistics.median(r['seconds'] for r in rows),
            'p95_seconds': sorted(r['seconds'] for r in rows)[math.ceil(.95 * len(rows)) - 1],
            'max_rss_mb': max(r['peak_rss_mb'] for r in rows),
        } for name, rows in report['cases'].items()}
        Path(args.output).write_text(json.dumps(report, indent=2), encoding='utf-8')
        print(json.dumps(report['summary'], indent=2))


if __name__ == '__main__':
    main()
